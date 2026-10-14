"""
Fill Performance report rows 3–4 (Excel rows 8–9, columns FG/FH) on MARCH26.
Does not change rows 1–2 of the metric list (on time + cancelled = rows 6–7).
"""
from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument(
        "input_xlsx",
        type=Path,
        help="Path to performance_excel.xlsx",
    )
    p.add_argument(
        "-o",
        "--output",
        type=Path,
        default=None,
        help="Output path (default: overwrite input)",
    )
    args = p.parse_args()
    in_path: Path = args.input_xlsx
    out_path: Path = args.output or in_path
    if out_path != in_path:
        shutil.copy2(in_path, out_path)
        work = out_path
    else:
        work = in_path

    wb = openpyxl.load_workbook(work)
    if "MARCH26" not in wb.sheetnames:
        print("Sheet MARCH26 not found. Available:", wb.sheetnames, file=sys.stderr)
        return 1
    sh = wb["MARCH26"]
    fg = column_index_from_string("FG")
    fh = column_index_from_string("FH")

    # Preserve metric rows 1–2 (rows 6–7)
    snap6 = (sh.cell(6, fg).value, sh.cell(6, fh).value)
    snap7 = (sh.cell(7, fg).value, sh.cell(7, fh).value)

    # Row 3: 3. Flight Operated Off–Sked; Row 4: a. Ground time saved
    sh.cell(8, fg).value = "=SUM(FG9:FG12)"
    sh.cell(8, fh).value = "=FG8/$FG$13"
    sh.cell(9, fg).value = "=EO5+EP5+EV5"
    sh.cell(9, fh).value = "=FG9/$FG$13"

    if (sh.cell(6, fg).value, sh.cell(6, fh).value) != snap6 or (
        sh.cell(7, fg).value,
        sh.cell(7, fh).value,
    ) != snap7:
        print("Error: rows 6–7 were modified.", file=sys.stderr)
        return 1

    wb.save(work)
    print(f"Updated {work}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
