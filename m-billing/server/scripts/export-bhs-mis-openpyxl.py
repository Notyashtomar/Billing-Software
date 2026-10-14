#!/usr/bin/env python3
"""Populate BHS MIS template with billing data — preserves Excel formulas (openpyxl)."""
from __future__ import annotations

import calendar
import json
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

RECO_THIN = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)
RECO_FONT = Font(name="Calibri", size=10, bold=False, color="000000")
RECO_FONT_BOLD = Font(name="Calibri", size=10, bold=True, color="000000")
RECO_FILL_WHITE = PatternFill(fill_type="solid", fgColor="FFFFFF")
RECO_FILL_HEADER = PatternFill(fill_type="solid", fgColor="D9E2F3")
RECO_FILL_TOTAL = PatternFill(fill_type="solid", fgColor="F2F2F2")
RECO_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=False)
RECO_ALIGN_HEADER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RECO_ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=False)


def is_domestic(category: str) -> bool:
    c = (category or "").lower()
    return "domestic" in c and "international" not in c


_TITLE_MONTH = re.compile(
    r"\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)'\s?\d{2}\b",
    re.I,
)


def title_month_stamp(from_ymd: str) -> str:
    """Match formatDelhiRevTitleMonth: 2026-08-01 → Aug' 26."""
    m = re.match(r"^(\d{4})-(\d{2})", from_ymd or "")
    if not m:
        return ""
    abbr = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    return f"{abbr[int(m.group(2)) - 1]}' {m.group(1)[2:]}"


def stamp_rev_summary_title(ws: Worksheet, from_ymd: str) -> None:
    stamp = title_month_stamp(from_ymd)
    if not stamp:
        return
    cell = ws["A1"]
    if type(cell).__name__ == "MergedCell" or not isinstance(cell.value, str):
        return
    text = _TITLE_MONTH.sub(stamp, cell.value, count=1)
    if text != cell.value:
        cell.value = text


def month_token(from_ymd: str) -> str:
    m = re.match(r"^(\d{4})-(\d{2})", from_ymd or "")
    if not m:
        return "MON-YY"
    abbr = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]
    return f"{abbr[int(m.group(2)) - 1]}-{m.group(1)[2:]}"


def set_value(ws: Worksheet, addr: str, value: Any) -> None:
    cell = ws[addr]
    if type(cell).__name__ == "MergedCell":
        return
    if cell.data_type == "f":
        return
    cell.value = value


def set_rc(ws: Worksheet, row: int, col: int, value: Any) -> None:
    cell = ws.cell(row, col)
    if type(cell).__name__ == "MergedCell":
        return
    if cell.data_type == "f":
        return
    cell.value = value


def set_rc_overwrite(ws: Worksheet, row: int, col: int, value: Any) -> None:
    cell = ws.cell(row, col)
    if type(cell).__name__ == "MergedCell":
        return
    cell.value = value


DEL_REV_COL = {
    "prevMonthFlights": 9,
    "invoiceAmount": 24,
    "invoiceAmountInr": 25,
    "prevInvoiceAmountInr": 28,
    "prevMonthOnly": 31,
}


BHS_REV_FIRST_ROW = 3
BHS_REV_COL = {
    "categoryDomInt": 2,
    "customerCode": 3,
    "airlineName": 4,
    "aircraftType": 5,
    "bodyType": 6,
    "flightCount": 7,
    "category": 8,
    "validFrom": 9,
    "validTo": 10,
    "currency": 15,
    "roe": 16,
    "packagePriceUsd": 17,
    "packagePriceInr": 18,
    "royaltyMode": 19,
    "invoiceTotal": 27,
    "invoiceNo": 29,
}
BHS_SCHED_FIRST_ROW = 2
BHS_SCHED_COL = {
    "flightNo": 1,
    "aircraftType": 2,
    "sta": 3,
    "std": 4,
    "dowName": 5,
    "body": 6,
    "dowNum": 7,
}
BHS_XXLD_FIRST_ROW = 26
BHS_XXLD_COL = {
    "serial": 1,
    "date": 2,
    "flightNo": 3,
    "aircraftType": 4,
    "customerCode": 15,
    "airlineName": 16,
    "qty": 17,
}


def replace_month_tokens(ws: Worksheet, token: str, station_label: str) -> None:
    is_chennai = (station_label or "").strip().lower() == "chennai"
    max_row = min(ws.max_row or 0, 6)
    max_col = min(ws.max_column or 0, 20)
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(r, c)
            if cell.data_type == "f" or not isinstance(cell.value, str):
                continue
            text = cell.value
            text = re.sub(r"[A-Z]{3}-\d{2}", token, text)
            text = re.sub(
                r"for the month of [A-Za-z]+-\d{2}",
                f"for the month of {token}",
                text,
                flags=re.I,
            )
            text = re.sub(
                r"for the m/o [A-Za-z]+-\d{2}",
                f"for the m/o {token}",
                text,
                flags=re.I,
            )
            if not is_chennai:
                text = re.sub(r"Chennai Stat\w*", f"{station_label} Station", text, flags=re.I)
                text = re.sub(
                    r"Bird Airport service\s*\(\s*Chennai\s*\)",
                    f"Bird Airport service({station_label})",
                    text,
                    flags=re.I,
                )
                text = re.sub(r"\bChennai\b", station_label, text, flags=re.I)
            if text != cell.value:
                cell.value = text


def populate_schedule_fris(ws: Worksheet, flights: list[dict[str, Any]]) -> None:
    unique: dict[str, dict[str, Any]] = {}
    for row in flights:
        key = f"{row.get('flightNo','')}|{row.get('aircraftType','')}"
        if key not in unique:
            unique[key] = row
    for r in range(2, 201):
        for c in range(1, 8):
            cell = ws.cell(r, c)
            if cell.data_type != "f":
                cell.value = None
    for idx, row in enumerate(unique.values()):
        r = BHS_SCHED_FIRST_ROW + idx
        set_rc(ws, r, BHS_SCHED_COL["flightNo"], row.get("flightNo"))
        set_rc(ws, r, BHS_SCHED_COL["aircraftType"], row.get("aircraftType"))
        set_rc(ws, r, BHS_SCHED_COL["sta"], row.get("sta") or "")
        set_rc(ws, r, BHS_SCHED_COL["std"], row.get("std") or "")
        sd = str(row.get("scheduleDate") or "")[:10]
        if sd:
            d = datetime.strptime(sd, "%Y-%m-%d")
            set_rc(ws, r, BHS_SCHED_COL["dowName"], d.strftime("%a"))
            set_rc(ws, r, BHS_SCHED_COL["dowNum"], d.isoweekday())
        cat = str(row.get("category") or "")
        set_rc(ws, r, BHS_SCHED_COL["body"], "WB" if "WB" in cat else "NB")


def _excel_col(n: int) -> str:
    out = ""
    while n > 0:
        n, rem = divmod(n - 1, 26)
        out = chr(65 + rem) + out
    return out


def _normalize_ymd(raw: Any) -> str:
    if isinstance(raw, datetime):
        return raw.strftime("%Y-%m-%d")
    text = str(raw or "").strip()
    if not text:
        return ""
    iso = re.match(r"^(\d{4})-(\d{2})-(\d{2})", text)
    if iso:
        return f"{iso.group(1)}-{iso.group(2)}-{iso.group(3)}"
    slash = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})", text)
    if slash:
        return f"{slash.group(3)}-{int(slash.group(2)):02d}-{int(slash.group(1)):02d}"
    dash = re.match(r"^(\d{1,2})-(\d{1,2})-(\d{4})$", text)
    if dash:
        return f"{dash.group(3)}-{int(dash.group(2)):02d}-{int(dash.group(1)):02d}"
    return ""


def _carrier_from_flight_no(flight_no: str) -> str:
    text = str(flight_no or "").strip().upper()
    if not text:
        return ""
    m = re.search(r"\b([A-Z0-9]{2})[-\s]*\d{1,5}\b", text)
    return m.group(1) if m else ""


_IATA_RE = re.compile(r"^(?:[A-Z]{2}|[A-Z][0-9]|[0-9][A-Z])$")
_NS_OPERATOR_CODES = {
    "ACS", "SKP", "VSR", "VSC", "AAPL", "ASPL", "FASP", "AAS", "ATS",
    "AEL", "AFO", "ASI", "FCS", "JSG", "TAL", "GVH", "GMR", "AIX",
}


def _is_scheduled_iata(code: str) -> bool:
    return bool(_IATA_RE.fullmatch(code))


def _excel_date(value: Any) -> Any:
    if isinstance(value, datetime):
        return value
    text = str(value or "")[:10]
    try:
        return datetime.strptime(text, "%Y-%m-%d")
    except ValueError:
        return value


def _royalty_inc_exc(value: Any) -> str:
    mode = str(value or "Inc").strip().lower()
    if mode in ("exc", "exclusive", "excluded"):
        return "Exc"
    return "Inc"


def _delhi_revenue_formula(row_idx: int) -> str:
    """Revenue excluding royalty. N/IFERROR stops text in the rate or flight cells becoming #VALUE!."""
    base = (
        f"ROUND(IFERROR(N(M{row_idx}),0)*(IFERROR(N(O{row_idx}),0)+IFERROR(N(P{row_idx}),0)),0)"
    )
    return (
        f'=IF(R{row_idx}="Exc",{base},'
        f'IF(B{row_idx}="Int",{base}*(100-$R$1)/100,{base}))'
    )


def _write_delhi_rate_inputs(ws: Worksheet, row_idx: int, row: dict[str, Any], roe: float) -> None:
    """Delhi S/U/V/W formulas add the USD and INR rates. Currency text must stay in column N."""
    currency = str(row.get("currency") or "INR").strip().upper() or "INR"
    usd = row.get("packagePriceUsd") or None
    inr = row.get("packagePriceInr") or None
    if currency == "USD":
        if usd is None and inr:
            usd = inr
        inr = None
    else:
        if inr is None and usd:
            inr = usd
        usd = None
    rate = row.get("roeToInr") if row.get("roeToInr") is not None else roe
    set_rc_overwrite(ws, row_idx, 14, currency)
    set_rc_overwrite(ws, row_idx, 15, float(usd) if usd else None)
    set_rc_overwrite(ws, row_idx, 16, float(inr) if inr else None)
    set_rc_overwrite(ws, row_idx, 17, float(rate) if rate else None)
    set_rc_overwrite(ws, row_idx, 18, _royalty_inc_exc(row.get("royaltyMode")))
    set_rc_overwrite(ws, row_idx, 10, int(row.get("currentFlightCount") or 0))
    set_rc_overwrite(ws, row_idx, 11, 0)
    set_rc_overwrite(ws, row_idx, 12, 0)
    set_rc_overwrite(ws, row_idx, 19, _delhi_revenue_formula(row_idx))
    set_rc_overwrite(
        ws,
        row_idx,
        23,
        f"=IF(IFERROR(N(O{row_idx}),0)>0,SUM(S{row_idx}:V{row_idx})*IFERROR(N(Q{row_idx}),1),SUM(S{row_idx}:V{row_idx}))",
    )


def _airline_full_name(airline_name: Any, customer_code: Any) -> str:
    code = str(customer_code or "").strip().upper().split("-")[0]
    name = re.sub(r"\s+", " ", str(airline_name or "")).strip()
    if not name:
        return ""
    prefixed = re.match(r"^([A-Za-z0-9]{2,4})\s*[-–/:]\s*(.+)$", name)
    if prefixed:
        prefix = prefixed.group(1).upper()
        rest = prefixed.group(2).strip()
        if (not code or prefix == code) and rest and rest.upper() != code:
            return rest
    if code and name.upper() == code:
        return ""
    return name


def _airline_code(customer_code: str, flight_no: str) -> str:
    from_flight = _carrier_from_flight_no(flight_no)
    if from_flight and _is_scheduled_iata(from_flight) and from_flight not in _NS_OPERATOR_CODES:
        return from_flight
    from_customer = str(customer_code or "").strip().upper().split("-")[0].strip()
    if (
        _is_scheduled_iata(from_customer)
        and from_customer not in _NS_OPERATOR_CODES
        and _carrier_from_flight_no(flight_no)
    ):
        return from_customer
    return "UNK"


def _is_non_scheduled(flight: dict[str, Any]) -> bool:
    blob = " ".join(
        str(flight.get(k) or "").lower()
        for k in ("category", "skdNonSkd", "ops")
    )
    if re.search(r"non[-\s]?scheduled|non[-\s]?skd|nonskd|nonschedule|charter", blob):
        return True
    if str(flight.get("bookingNo") or "").strip() or str(flight.get("challanNo") or "").strip():
        return True
    code = str(flight.get("customerCode") or "").strip().upper()
    if code in _NS_OPERATOR_CODES:
        return True
    return not bool(_carrier_from_flight_no(str(flight.get("flightNo") or "")))


def _today_ist_ymd() -> str:
    # Asia/Kolkata = UTC+5:30
    from datetime import timedelta, timezone

    ist = timezone(timedelta(hours=5, minutes=30))
    return datetime.now(ist).strftime("%Y-%m-%d")


def _proj_layout(kind: str) -> dict[str, Any]:
    if kind == "del":
        return {
            "first_airline": 4,
            "weekday": 3,
            "day_name": 2,
            "protect_from": 67,
            "total_header": "Total\n(A)",
        }
    return {
        "first_airline": 3,
        "weekday": 2,
        "day_name": None,
        "protect_from": 25,
        "total_header": "Total",
    }


def build_proj_skd_matrix(
    flights: list[dict[str, Any]],
    from_date_str: str | None,
    to_date_str: str | None = None,
) -> dict[str, Any]:
    target = (from_date_str or "")[:10]
    try:
        dt_start = datetime.strptime(target, "%Y-%m-%d")
        year, month = dt_start.year, dt_start.month
    except Exception:
        year, month = datetime.now().year, datetime.now().month
    _, num_days = calendar.monthrange(year, month)
    month_start = f"{year:04d}-{month:02d}-01"
    month_end = f"{year:04d}-{month:02d}-{num_days:02d}"
    range_end = (to_date_str or month_end)[:10] or month_end
    as_of = min(range_end, month_end, _today_ist_ymd())

    columns: dict[str, dict[str, str]] = {}
    counts: dict[str, dict[str, int]] = {}
    seen: set[str] = set()

    for f in flights:
        if f.get("cancelled"):
            continue
        if _is_non_scheduled(f):
            continue
        day = _normalize_ymd(f.get("scheduleDate"))
        if not day or day < month_start or day > as_of:
            continue
        airline = _airline_code(str(f.get("customerCode") or ""), str(f.get("flightNo") or ""))
        if not airline or airline == "UNK":
            continue
        label = re.sub(r"\s+", " ", str(f.get("aircraftType") or "").strip())
        ac_key = label.upper()
        col_key = f"{airline}|{ac_key}"
        if col_key not in columns:
            columns[col_key] = {
                "key": col_key,
                "airlineCode": airline,
                "aircraftType": ac_key,
                "aircraftLabel": label,
            }
        flight_no = str(f.get("flightNo") or "").strip().upper() or "UNK"
        dedupe = f"{day}|{flight_no}|{col_key}"
        if dedupe in seen:
            continue
        seen.add(dedupe)
        counts.setdefault(day, {})
        counts[day][col_key] = counts[day].get(col_key, 0) + 1

    col_list = sorted(columns.values(), key=lambda c: (c["airlineCode"], c["aircraftType"]))
    return {
        "year": year,
        "month": month,
        "numDays": num_days,
        "asOfYmd": as_of,
        "columns": col_list,
        "counts": counts,
    }


def _proj_clear_until(num_days: int, template_max_row: int) -> int:
    total_row = 3 + max(num_days, 0) + 1
    return max(total_row + 6, template_max_row, 40)


def _unmerge_proj_table(ws: Worksheet, protect: int) -> None:
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row < 2:
            continue
        if rng.min_col >= protect:
            continue
        if rng.min_row == 2 and rng.max_row == 2 and rng.min_col <= 2 and rng.max_col <= 2:
            continue
        ws.unmerge_cells(str(rng))


def _split_projected(columns: list[dict[str, Any]], reco_rows: list[dict[str, Any]]) -> list[int]:
    nets: dict[str, int] = {}
    for row in reco_rows:
        code = str(row.get("airlineCode") or "").strip().upper()
        if not code:
            continue
        nets[code] = nets.get(code, 0) + int(row.get("netCurrent") or 0)
    counts: dict[str, int] = {}
    for col in columns:
        code = str(col.get("airlineCode") or "").strip().upper()
        counts[code] = counts.get(code, 0) + 1
    leftover = dict(nets)
    seen: dict[str, int] = {}
    out: list[int] = []
    for col in columns:
        code = str(col.get("airlineCode") or "").strip().upper()
        seen[code] = seen.get(code, 0) + 1
        n = counts.get(code, 1)
        total = nets.get(code, 0)
        if seen[code] < n:
            part = total // n
            leftover[code] = leftover.get(code, 0) - part
            out.append(part)
        else:
            out.append(leftover.get(code, 0))
    return out


def populate_proj_skd(
    ws: Worksheet,
    flights: list[dict[str, Any]],
    token: str,
    from_date_str: str | None = None,
    to_date_str: str | None = None,
    layout_kind: str = "maa",
    reco: dict[str, Any] | None = None,
) -> dict[str, Any]:
    matrix = build_proj_skd_matrix(flights, from_date_str, to_date_str)
    layout = _proj_layout(layout_kind)
    first_ac = layout["first_airline"]
    protect = layout["protect_from"]
    max_cols = max(1, protect - first_ac - 1)
    columns = matrix["columns"][:max_cols]
    last_airline = first_ac + len(columns) - 1 if columns else first_ac
    total_col = last_airline + 1 if columns else first_ac
    first_day_row = 4
    last_day_row = 3 + matrix["numDays"]
    total_row = last_day_row + 1
    projected_row = total_row + 1
    diff_row = total_row + 2
    first_letter = _excel_col(first_ac)
    last_letter = _excel_col(max(first_ac, last_airline))
    last_write = min(protect - 1, max(total_col, first_ac))

    _unmerge_proj_table(ws, protect)
    if ws["A2"].data_type != "f":
        ws["A2"] = f"Flight Chart for the m/o {token}"
    ws["A2"].font = Font(name="Calibri", size=11, bold=True, color="000000")

    clear_until = _proj_clear_until(matrix["numDays"], int(ws.max_row or 0))
    for r in range(2, clear_until + 1):
        ws.row_dimensions[r].height = 28 if r == 2 else 16
        for c in range(1, protect):
            if r == 2 and c <= 2:
                continue
            cell = ws.cell(r, c)
            cell.value = None
            _style_reco_cell(cell, "plain")
            cell.number_format = "General"

    ws.cell(3, 1).value = "Date "
    if layout["day_name"]:
        ws.cell(3, layout["day_name"]).value = "Days"
    ws.cell(3, layout["weekday"]).value = "Weekdays"

    for idx, col in enumerate(columns):
        c = first_ac + idx
        label = re.sub(r"\s+", " ", str(col["aircraftLabel"] or col["aircraftType"] or "").strip())
        ws.cell(2, c).value = label
        ws.cell(3, c).value = col["airlineCode"]
        ws.column_dimensions[_excel_col(c)].width = max(11, min(16, len(label or col["airlineCode"]) + 2))
    ws.cell(3, total_col).value = str(layout["total_header"]).replace("\n", " ")
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions[_excel_col(layout["weekday"])].width = 10
    if layout["day_name"]:
        ws.column_dimensions[_excel_col(layout["day_name"])].width = 8
    ws.column_dimensions[_excel_col(total_col)].width = 10

    for c in range(1, last_write + 1):
        _style_reco_cell(ws.cell(2, c), "header")
        _style_reco_cell(ws.cell(3, c), "header")
        ws.cell(2, c).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    _style_reco_cell(ws.cell(2, 1), "plain")
    ws.cell(2, 1).font = Font(name="Calibri", size=11, bold=True, color="000000")

    def paint(r: int, kind: str) -> None:
        for c in range(1, last_write + 1):
            cell = ws.cell(r, c)
            _style_reco_cell(cell, kind)
            if c == 1 and kind == "data":
                cell.number_format = "YYYY-MM-DD"
            elif c >= layout["weekday"]:
                cell.number_format = "0"

    year, month = matrix["year"], matrix["month"]
    for day_num in range(1, matrix["numDays"] + 1):
        r = 3 + day_num
        day_str = f"{year:04d}-{month:02d}-{day_num:02d}"
        if day_num == 1:
            ws.cell(r, 1).value = datetime(year, month, day_num)
        else:
            ws.cell(r, 1).value = f"=A{r-1}+1"
        if layout["day_name"]:
            ws.cell(r, layout["day_name"]).value = f'=TEXT(A{r},"DDD")'
        ws.cell(r, layout["weekday"]).value = f"=WEEKDAY(A{r},2)"
        for idx, col in enumerate(columns):
            ws.cell(r, first_ac + idx).value = matrix["counts"].get(day_str, {}).get(col["key"], 0)
        if columns:
            ws.cell(r, total_col).value = f"=SUM({first_letter}{r}:{last_letter}{r})"
        else:
            ws.cell(r, total_col).value = 0
        paint(r, "data")

    ws.cell(total_row, 1).value = "Total"
    for idx, _col in enumerate(columns):
        letter = _excel_col(first_ac + idx)
        ws.cell(total_row, first_ac + idx).value = f"=SUM({letter}{first_day_row}:{letter}{last_day_row})"
    tot_letter = _excel_col(total_col)
    ws.cell(total_row, total_col).value = f"=SUM({tot_letter}{first_day_row}:{tot_letter}{last_day_row})"
    paint(total_row, "total")

    projected = _split_projected(columns, list((reco or {}).get("rows") or []))
    ws.cell(projected_row, 1).value = "Projected Schedule"
    for idx, value in enumerate(projected):
        ws.cell(projected_row, first_ac + idx).value = value
    if columns:
        ws.cell(projected_row, total_col).value = (
            f"=SUM({first_letter}{projected_row}:{last_letter}{projected_row})"
        )
    paint(projected_row, "data")
    ws.cell(projected_row, 1).alignment = RECO_ALIGN_LEFT

    ws.cell(diff_row, 1).value = "Difference"
    start = first_ac if columns else total_col
    for c in range(start, total_col + 1):
        letter = _excel_col(c)
        ws.cell(diff_row, c).value = f"={letter}{total_row}-{letter}{projected_row}"
    paint(diff_row, "data")
    ws.cell(diff_row, 1).alignment = RECO_ALIGN_LEFT

    matrix["columns"] = columns
    return matrix


def _reco_layout(layout_kind: str) -> dict[str, Any]:
    if layout_kind == "del":
        return {
            "series": 1,
            "code": 2,
            "flight": None,
            "first_wd": 3,
            "weekly": 10,
            "ops": 11,
            "completed": 12,
            "extra0": 13,
            "extra_n": 3,
            "total": 16,
            "add": 17,
            "net_cm": 18,
            "net_lm": 19,
            "diff": 20,
            "reason": 21,
            "last_col": 21,
        }
    return {
        "series": 1,
        "code": 2,
        "flight": 3,
        "first_wd": 4,
        "weekly": 11,
        "ops": 12,
        "completed": 13,
        "extra0": 14,
        "extra_n": 3,
        "total": 17,
        "add": 18,
        "net_cm": 19,
        "net_lm": 20,
        "diff": 21,
        "reason": 22,
        "last_col": 22,
    }


def _reco_clear_until(row_count: int, template_max_row: int) -> int:
    needed = 4 + max(row_count, 0) + 6 + 8
    return max(needed, template_max_row, 40)


def _unmerge_reco_table(ws: Worksheet) -> None:
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= 3:
            ws.unmerge_cells(str(rng))


def _style_reco_cell(cell: Any, kind: str) -> None:
    if kind == "header":
        cell.font = RECO_FONT_BOLD
        cell.fill = RECO_FILL_HEADER
        cell.alignment = RECO_ALIGN_HEADER
        cell.border = RECO_THIN
    elif kind == "total":
        cell.font = RECO_FONT_BOLD
        cell.fill = RECO_FILL_TOTAL
        cell.alignment = RECO_ALIGN
        cell.border = RECO_THIN
    elif kind == "data":
        cell.font = RECO_FONT
        cell.fill = RECO_FILL_WHITE
        cell.alignment = RECO_ALIGN
        cell.border = RECO_THIN
    else:
        cell.font = RECO_FONT
        cell.fill = RECO_FILL_WHITE
        cell.alignment = RECO_ALIGN_LEFT
        cell.border = Border()


def populate_reco_skd(ws: Worksheet, reco: dict[str, Any], layout_kind: str) -> None:
    layout = _reco_layout(layout_kind)
    rows = reco.get("rows") or []
    extras = reco.get("extraDays") or []
    data_start = 4
    _unmerge_reco_table(ws)
    clear_until = _reco_clear_until(len(rows), int(ws.max_row or 0))
    last_col = layout["last_col"]
    for r in range(data_start, clear_until + 1):
        ws.row_dimensions[r].height = 16
        for c in range(1, last_col + 1):
            cell = ws.cell(r, c)
            cell.value = None
            _style_reco_cell(cell, "plain")
            cell.number_format = "General"

    for slot in range(layout["extra_n"]):
        extra = extras[slot] if slot < len(extras) else None
        ws.cell(3, layout["extra0"] + slot).value = (
            f"Extra day-{extra['weekday']}" if extra else ""
        )
    ws.cell(3, layout["net_cm"]).value = f"Net Total Flights {reco.get('monthToken') or ''}"
    ws.cell(3, layout["net_lm"]).value = f"Net Total Flights {reco.get('prevMonthToken') or ''}"
    for c in range(1, last_col + 1):
        _style_reco_cell(ws.cell(3, c), "header")

    numeric_cols = set(
        [layout["series"]]
        + list(range(layout["first_wd"], layout["first_wd"] + 7))
        + [layout["weekly"], layout["ops"], layout["completed"]]
        + list(range(layout["extra0"], layout["extra0"] + layout["extra_n"]))
        + [layout["total"], layout["add"], layout["net_cm"], layout["net_lm"], layout["diff"]]
    )

    def paint_row(r: int, kind: str) -> None:
        for c in range(1, last_col + 1):
            cell = ws.cell(r, c)
            _style_reco_cell(cell, kind)
            if c in numeric_cols:
                cell.number_format = "0"

    for idx, row in enumerate(rows):
        r = data_start + idx
        ws.cell(r, layout["series"]).value = idx + 1
        ws.cell(r, layout["code"]).value = row.get("airlineCode")
        if layout["flight"] is not None:
            ws.cell(r, layout["flight"]).value = row.get("flightNo")
        weekday_counts = row.get("weekdayCounts") or [0] * 7
        for wd in range(7):
            count = weekday_counts[wd] if wd < len(weekday_counts) else 0
            ws.cell(r, layout["first_wd"] + wd).value = int(count or 0)
        ws.cell(r, layout["weekly"]).value = int(row.get("weeklyTotal") or 0)
        ws.cell(r, layout["ops"]).value = int(row.get("fullOpsWeeks") or 4)
        ws.cell(r, layout["completed"]).value = int(row.get("completedWeeks") or 0)
        extra_counts = row.get("extraCounts") or []
        for slot in range(layout["extra_n"]):
            ws.cell(r, layout["extra0"] + slot).value = (
                int(extra_counts[slot]) if slot < len(extra_counts) else 0
            )
        extra_sum = int(row.get("extraSum") or 0)
        completed = int(row.get("completedWeeks") or 0)
        ws.cell(r, layout["total"]).value = completed + extra_sum
        ws.cell(r, layout["add"]).value = int(row.get("addLess") or 0)
        ws.cell(r, layout["net_cm"]).value = int(row.get("netCurrent") or 0)
        ws.cell(r, layout["net_lm"]).value = int(row.get("netLast") or 0)
        ws.cell(r, layout["diff"]).value = int(row.get("diff") or 0)
        ws.cell(r, layout["reason"]).value = row.get("reason") or ""
        paint_row(r, "data")

    first = data_start
    last = data_start + max(len(rows), 1) - 1
    tot = data_start + len(rows)
    label_col = layout["flight"] if layout["flight"] is not None else layout["code"]
    ws.cell(tot, label_col).value = "Total"
    if rows:
        sum_cols = (
            list(range(layout["first_wd"], layout["first_wd"] + 7))
            + [layout["weekly"], layout["completed"]]
            + list(range(layout["extra0"], layout["extra0"] + layout["extra_n"]))
            + [layout["total"], layout["add"], layout["net_cm"], layout["net_lm"], layout["diff"]]
        )
        for c in sum_cols:
            letter = _excel_col(c)
            ws.cell(tot, c).value = f"=SUM({letter}{first}:{letter}{last})"
    paint_row(tot, "total")

    if layout_kind != "del":
        reco_r = tot + 1
        ws.cell(reco_r, 9).value = "Flight Reco:"
        ws.cell(reco_r + 1, 9).value = f"Weekly Flight of {reco.get('prevMonthToken') or ''}"
        ws.cell(reco_r + 1, 11).value = reco.get("weeklyLast") or 0
        ws.cell(reco_r + 2, 9).value = f"Weekly Flight of {reco.get('monthToken') or ''}"
        ws.cell(reco_r + 2, 11).value = f"=K{tot}"
        ws.cell(reco_r + 3, 9).value = "Difference:-"
        ws.cell(reco_r + 3, 11).value = f"=K{reco_r + 1}-K{reco_r + 2}"
        for r in range(reco_r, reco_r + 4):
            _style_reco_cell(ws.cell(r, 9), "plain")
            _style_reco_cell(ws.cell(r, 11), "plain")
            ws.cell(r, 11).number_format = "0"
        ws.cell(reco_r, 9).font = RECO_FONT_BOLD


def populate_rev_summary(
    ws: Worksheet,
    price_rows: list[dict[str, Any]],
    invoices: list[dict[str, Any]],
    roe: float,
    flights: list[dict[str, Any]] | None = None,
    rev_rows: list[dict[str, Any]] | None = None,
    rev_headers: dict[str, Any] | None = None,
) -> None:
    inv_map: dict[str, dict[str, Any]] = {}
    for inv in invoices:
        code = str(inv.get("customerCode") or "").upper()
        if code:
            base = code.split("-")[0]
            if code not in inv_map:
                inv_map[code] = inv
            if base not in inv_map:
                inv_map[base] = inv
        bill_to = str(inv.get("billToCustomerCode") or "").upper()
        if bill_to:
            base_b = bill_to.split("-")[0]
            if bill_to not in inv_map:
                inv_map[bill_to] = inv
            if base_b not in inv_map:
                inv_map[base_b] = inv

    by_airline_count: dict[str, int] = {}
    by_airline_ac_count: dict[str, int] = {}
    for f in flights or []:
        # Never keep MAA template flight counts — login station data only.
        if f.get("cancelled") or f.get("eligible") is not True:
            continue
        f_code = str(f.get("customerCode") or "").upper().split("-")[0]
        f_ac = str(f.get("aircraftType") or "").upper().strip()
        by_airline_count[f_code] = by_airline_count.get(f_code, 0) + 1
        if f_ac:
            key = f"{f_code}|{f_ac}"
            by_airline_ac_count[key] = by_airline_ac_count.get(key, 0) + 1

    start_r = BHS_REV_FIRST_ROW
    indexed = rev_rows if rev_rows else None
    if indexed is None:
        indexed = []
        for pc in price_rows:
            code = str(pc.get("customerCode") or "").upper()
            base_code = code.split("-")[0] if "-" in code else code
            inv = inv_map.get(code) or inv_map.get(base_code)
            ac_type = str(pc.get("aircraftType") or "").upper().strip()
            flt_count = by_airline_ac_count.get(f"{base_code}|{ac_type}")
            if flt_count is None or flt_count == 0:
                flt_count = by_airline_count.get(base_code, 0)
            indexed.append(
                {
                    "categoryDomInt": "Dom" if is_domestic(str(pc.get("category") or "")) else "Int",
                    "customerCode": base_code,
                    "airlineName": pc.get("airlineName"),
                    "aircraftType": pc.get("aircraftType"),
                    "bodyType": pc.get("bodyType") or "NB",
                    "flightCount": flt_count,
                    "category": pc.get("category"),
                    "validFrom": pc.get("validFrom"),
                    "validTo": pc.get("validTo"),
                    "currency": pc.get("currency") or "INR",
                    "roeToInr": roe,
                    "packagePriceUsd": pc.get("packagePriceUsd"),
                    "packagePriceInr": pc.get("packagePriceInr"),
                    "royaltyMode": pc.get("royaltyMode"),
                    "invoiceTotal": (inv or {}).get("totalInclTax"),
                    "invoiceNumber": (inv or {}).get("invoiceNumber"),
                    "prevMonthFlightCount": by_airline_count.get(base_code, 0),
                    "currentFlightCount": flt_count,
                    "invoiceAmount": (inv or {}).get("totalInclTax") or 0,
                    "invoiceAmountInr": (inv or {}).get("totalInclTax") or 0,
                    "prevInvoiceAmountInr": 0,
                }
            )

    is_maa_sheet = ws.title == "Rev_summary" or (ws["J51"].value == "Type" or ws["J52"].value == "NB")
    headers = rev_headers or {}
    if not is_maa_sheet and headers:
        set_rc_overwrite(ws, 2, DEL_REV_COL["prevMonthFlights"], headers.get("flightOf"))
        set_rc_overwrite(ws, 2, DEL_REV_COL["invoiceAmount"], headers.get("invoiceAmount"))
        set_rc_overwrite(ws, 2, DEL_REV_COL["invoiceAmountInr"], headers.get("invoiceInr"))
        set_rc_overwrite(ws, 2, DEL_REV_COL["prevInvoiceAmountInr"], headers.get("prevInr"))
        set_rc_overwrite(ws, 2, DEL_REV_COL["prevMonthOnly"], headers.get("prevOnly"))

    seen_invoice_airlines: set[str] = set()
    for idx, row in enumerate(indexed):
        r = start_r + idx
        airline_name = _airline_full_name(row.get("airlineName"), row.get("customerCode"))
        set_rc(ws, r, BHS_REV_COL["categoryDomInt"], row.get("categoryDomInt"))
        if is_maa_sheet:
            set_rc(ws, r, BHS_REV_COL["customerCode"], row.get("customerCode"))
            set_rc(ws, r, BHS_REV_COL["airlineName"], airline_name)
        else:
            set_rc_overwrite(ws, r, BHS_REV_COL["customerCode"], airline_name)
            set_rc(ws, r, BHS_REV_COL["airlineName"], row.get("bodyType") or "NB")
        set_rc(ws, r, BHS_REV_COL["aircraftType"], row.get("aircraftType"))
        set_rc(ws, r, BHS_REV_COL["bodyType"], row.get("bodyType") or "NB")
        if is_maa_sheet:
            set_rc(ws, r, BHS_REV_COL["flightCount"], row.get("flightCount") or 0)
        elif row.get("validFrom"):
            cell = ws.cell(r, BHS_REV_COL["flightCount"])
            if type(cell).__name__ != "MergedCell":
                cell.value = _excel_date(row.get("validFrom"))
                cell.number_format = "MM-DD-YY"
        if is_maa_sheet:
            set_rc(ws, r, BHS_REV_COL["category"], row.get("category"))
        if is_maa_sheet and row.get("validFrom"):
            set_rc(ws, r, BHS_REV_COL["validFrom"], row.get("validFrom"))
        if is_maa_sheet and row.get("validTo"):
            set_rc(ws, r, BHS_REV_COL["validTo"], row.get("validTo"))
        if is_maa_sheet:
            set_rc(ws, r, BHS_REV_COL["currency"], row.get("currency") or "INR")
            set_rc(ws, r, BHS_REV_COL["roe"], row.get("roeToInr") if row.get("roeToInr") is not None else roe)
            if row.get("packagePriceUsd"):
                set_rc(ws, r, BHS_REV_COL["packagePriceUsd"], row.get("packagePriceUsd"))
            if row.get("packagePriceInr"):
                set_rc(ws, r, BHS_REV_COL["packagePriceInr"], row.get("packagePriceInr"))
            if row.get("royaltyMode"):
                set_rc(ws, r, BHS_REV_COL["royaltyMode"], row.get("royaltyMode"))
        else:
            _write_delhi_rate_inputs(ws, r, row, roe)
        if is_maa_sheet and row.get("invoiceTotal") is not None:
            set_rc(ws, r, BHS_REV_COL["invoiceTotal"], row.get("invoiceTotal"))
        if is_maa_sheet and row.get("invoiceNumber"):
            set_rc(ws, r, BHS_REV_COL["invoiceNo"], row.get("invoiceNumber"))
        if not is_maa_sheet:
            set_rc_overwrite(ws, r, DEL_REV_COL["prevMonthFlights"], row.get("prevMonthFlightCount") or 0)
            airline = str(row.get("customerCode") or "").upper().split("-")[0]
            first_airline_row = bool(airline) and airline not in seen_invoice_airlines
            if first_airline_row:
                seen_invoice_airlines.add(airline)
            invoice_amount = row.get("invoiceAmount") or 0 if first_airline_row else 0
            invoice_inr = row.get("invoiceAmountInr") or 0 if first_airline_row else 0
            prev_inr = row.get("prevInvoiceAmountInr") or 0 if first_airline_row else 0
            set_rc_overwrite(ws, r, DEL_REV_COL["invoiceAmount"], invoice_amount)
            set_rc_overwrite(ws, r, DEL_REV_COL["invoiceAmountInr"], invoice_inr)
            set_rc_overwrite(ws, r, DEL_REV_COL["prevInvoiceAmountInr"], prev_inr)
            set_rc_overwrite(ws, r, DEL_REV_COL["prevMonthOnly"], prev_inr)

    end_data_row = max(start_r + len(indexed) - 1, start_r)
    is_maa = ws.title == "Rev_summary" or (ws["J51"].value == "Type" or ws["J52"].value == "NB")

    # Clear template placeholder rows beyond end_data_row
    for clear_r in range(end_data_row + 1, 65):
        for col_idx in range(1, 40):
            cell = ws.cell(clear_r, col_idx)
            if type(cell).__name__ != "MergedCell":
                cell.value = None

    total_row = end_data_row + 1

    if not is_maa:
        # Standard / DEL Layout: place summary table dynamically after main rows
        set_value(ws, f"D{total_row}", "Total")
        set_value(ws, f"T{total_row}", f"=SUM(T{start_r}:T{end_data_row})")
        set_value(ws, f"U{total_row}", f"=SUM(U{start_r}:U{end_data_row})")
        set_value(ws, f"V{total_row}", f"=SUM(V{start_r}:V{end_data_row})")
        set_value(ws, f"W{total_row}", f"=SUM(W{start_r}:W{end_data_row})")
        set_value(ws, f"X{total_row}", f"=SUM(X{start_r}:X{end_data_row})")
        set_value(ws, f"Y{total_row}", f"=SUM(Y{start_r}:Y{end_data_row})")
        set_value(ws, f"Z{total_row}", f"=SUM(Z{start_r}:Z{end_data_row})")
        set_value(ws, f"AA{total_row}", f"=SUM(AA{start_r}:AA{end_data_row})")
        set_value(ws, f"AB{total_row}", f"=SUM(AB{start_r}:AB{end_data_row})")

        summary_start = total_row + 2

        set_value(ws, f"H{summary_start}", "Per Flight")
        set_value(ws, f"I{summary_start}", "Revenue [Including GST]")
        set_value(ws, f"J{summary_start}", "%")
        set_value(ws, f"K{summary_start}", "# of Flt in CM")
        set_value(ws, f"L{summary_start}", "INR")
        set_value(ws, f"M{summary_start}", f"=AB{total_row}/SUM(M{start_r}:M{end_data_row})")

        s_dom = summary_start + 1
        s_int = summary_start + 2
        s_tot = summary_start + 3

        ws[f"H{s_dom}"].number_format = "General"
        set_value(ws, f"H{s_dom}", 46446)
        set_value(ws, f"I{s_dom}", f'=SUMIF(B{start_r}:B{end_data_row},"Dom",Z{start_r}:Z{end_data_row})')
        set_value(ws, f"J{s_dom}", f"=I{s_dom}/I{s_tot}")
        set_value(ws, f"K{s_dom}", f'=SUMIF(B{start_r}:B{end_data_row},"Dom",M{start_r}:M{end_data_row})')
        set_value(ws, f"L{s_dom}", f"=IF(K{s_dom}>0,I{s_dom}/K{s_dom},0)")
        set_value(ws, f"M{s_dom}", f"=L{s_dom}-H{s_dom}")

        ws[f"H{s_int}"].number_format = "General"
        set_value(ws, f"H{s_int}", 46446)
        set_value(ws, f"I{s_int}", f'=SUMIF(B{start_r}:B{end_data_row},"Int",Z{start_r}:Z{end_data_row})')
        set_value(ws, f"J{s_int}", f"=I{s_int}/I{s_tot}")
        set_value(ws, f"K{s_int}", f'=SUMIF(B{start_r}:B{end_data_row},"Int",M{start_r}:M{end_data_row})')
        set_value(ws, f"L{s_int}", f"=IF(K{s_int}>0,I{s_int}/K{s_int},0)")
        set_value(ws, f"M{s_int}", f"=L{s_int}-H{s_int}")

        set_value(ws, f"I{s_tot}", f"=SUM(I{s_dom}:I{s_int})")
        set_value(ws, f"J{s_tot}", f"=SUM(J{s_dom}:J{s_int})")
        set_value(ws, f"K{s_tot}", f"=SUM(K{s_dom}:K{s_int})")
        set_value(ws, f"L{s_tot}", f"=IF(K{s_tot}>0,I{s_tot}/K{s_tot},0)")
        set_value(ws, f"M{s_tot}", f"=L{s_tot}-H{s_int}")

    else:
        # MAA Layout: place summary table dynamically after main rows
        set_value(ws, f"G{total_row}", f"=SUM(G{start_r}:G{end_data_row})")
        set_value(ws, f"I{total_row}", "Total Flights")
        set_value(ws, f"K{total_row}", f"=SUM(K{start_r}:K{end_data_row})")
        set_value(ws, f"L{total_row}", f"=SUM(L{start_r}:L{end_data_row})")
        set_value(ws, f"M{total_row}", f"=SUM(M{start_r}:M{end_data_row})")
        set_value(ws, f"N{total_row}", f"=SUM(N{start_r}:N{end_data_row})")
        set_value(ws, f"T{total_row}", f"=SUM(T{start_r}:T{end_data_row})")
        set_value(ws, f"Z{total_row}", f"=SUM(Z{start_r}:Z{end_data_row})")

        summary_start = total_row + 3
        set_value(ws, f"J{summary_start}", "Type")
        set_value(ws, f"K{summary_start}", "Per Flight")
        set_value(ws, f"L{summary_start}", "Revenue [Including GST]")
        set_value(ws, f"M{summary_start}", "%")
        set_value(ws, f"N{summary_start}", "# of Flt in CM")
        set_value(ws, f"O{summary_start}", "# of Flt in LM")
        set_value(ws, f"P{summary_start}", "Difference")

        s_nb = summary_start + 1
        s_wb = summary_start + 2
        s_tot = summary_start + 3
        s_diff = summary_start + 4

        set_value(ws, f"J{s_nb}", "NB")
        set_value(ws, f"K{s_nb}", f"=IF(N{s_nb}>0,L{s_nb}/N{s_nb},0)")
        set_value(ws, f"L{s_nb}", f'=SUMIF($F${start_r}:$F${end_data_row},J{s_nb},$Z${start_r}:$Z${end_data_row})')
        set_value(ws, f"M{s_nb}", f"=L{s_nb}/L{s_tot}")
        set_value(ws, f"N{s_nb}", f'=SUMIF($F${start_r}:$F${end_data_row},J{s_nb},$N${start_r}:$N${end_data_row})')
        set_value(ws, f"O{s_nb}", f'=SUMIF($F${start_r}:$F${end_data_row},J{s_nb},$G${start_r}:$G${end_data_row})')
        set_value(ws, f"P{s_nb}", f"=N{s_nb}-O{s_nb}")

        set_value(ws, f"J{s_wb}", "WB")
        set_value(ws, f"K{s_wb}", f"=IF(N{s_wb}>0,L{s_wb}/N{s_wb},0)")
        set_value(ws, f"L{s_wb}", f'=SUMIF($F${start_r}:$F${end_data_row},J{s_wb},$Z${start_r}:$Z${end_data_row})')
        set_value(ws, f"M{s_wb}", f"=L{s_wb}/L{s_tot}")
        set_value(ws, f"N{s_wb}", f'=SUMIF($F${start_r}:$F${end_data_row},J{s_wb},$N${start_r}:$N${end_data_row})')
        set_value(ws, f"O{s_wb}", f'=SUMIF($F${start_r}:$F${end_data_row},J{s_wb},$G${start_r}:$G${end_data_row})')
        set_value(ws, f"P{s_wb}", f"=N{s_wb}-O{s_wb}")

        set_value(ws, f"J{s_tot}", "Total")
        set_value(ws, f"L{s_tot}", f"=SUM(L{s_nb}:L{s_wb})")
        set_value(ws, f"M{s_tot}", f"=SUM(M{s_nb}:M{s_wb})")
        set_value(ws, f"N{s_tot}", f"=SUM(N{s_nb}:N{s_wb})")
        set_value(ws, f"O{s_tot}", f"=SUM(O{s_nb}:O{s_wb})")
        set_value(ws, f"P{s_tot}", f"=SUM(P{s_nb}:P{s_wb})")

        set_value(ws, f"J{s_diff}", "Difference")
        set_value(ws, f"L{s_diff}", f"=L{s_tot}-Z{total_row}")
        set_value(ws, f"N{s_diff}", f"=N{s_tot}-N{total_row}")
        set_value(ws, f"O{s_diff}", f"=O{s_tot}-G{total_row}")


def populate_addl_xxld(
    ws: Worksheet,
    flights: list[dict[str, Any]],
    token: str,
    station_label: str,
) -> None:
    cancelled = [f for f in flights if f.get("cancelled")]
    ws["A1"] = f"Bird Airport service({station_label}) Private Limited"
    ws["A2"] = f"Details of cancelled flights in {station_label} for the month of {token}"
    for idx, row in enumerate(cancelled[:100]):
        r = BHS_XXLD_FIRST_ROW + idx
        code = str(row.get("customerCode") or "").upper().split("-")[0]
        set_rc(ws, r, BHS_XXLD_COL["serial"], idx + 1)
        set_rc(ws, r, BHS_XXLD_COL["date"], row.get("scheduleDate"))
        set_rc(ws, r, BHS_XXLD_COL["flightNo"], row.get("flightNo"))
        set_rc(ws, r, BHS_XXLD_COL["aircraftType"], row.get("aircraftType"))
        set_rc(ws, r, BHS_XXLD_COL["customerCode"], code)
        set_rc(ws, r, BHS_XXLD_COL["airlineName"], row.get("airlineName"))
        set_rc(ws, r, BHS_XXLD_COL["qty"], 1)


def export(payload: dict[str, Any]) -> None:
    template = Path(payload["templatePath"])
    output = Path(payload["outputPath"])
    profile = payload["profile"]
    token = month_token(payload["from"])
    station_label = payload.get("stationLabel") or payload.get("stationCode", "STN")

    wb = openpyxl.load_workbook(template)
    for sheet_def in profile.get("sheets", []):
        name = sheet_def.get("sheetName")
        if name not in wb.sheetnames:
            continue
        replace_month_tokens(wb[name], token, station_label)

    flights = payload.get("flights") or []
    if profile.get("scheduleSheet") and profile["scheduleSheet"] in wb.sheetnames:
        populate_schedule_fris(wb[profile["scheduleSheet"]], flights)
    proj_matrix = None
    if profile.get("projSheet") in wb.sheetnames:
        proj_matrix = populate_proj_skd(
            wb[profile["projSheet"]],
            flights,
            token,
            payload.get("from"),
            payload.get("to"),
            str(profile.get("layout") or "maa"),
            payload.get("reco") or {},
        )
    reco_name = profile.get("recoSheet")
    if reco_name in wb.sheetnames:
        populate_reco_skd(
            wb[reco_name],
            payload.get("reco") or {},
            str(profile.get("layout") or "maa"),
        )
    rev = profile.get("revSummarySheet")
    if rev in wb.sheetnames:
        populate_rev_summary(
            wb[rev],
            payload.get("priceCardRows") or [],
            payload.get("invoices") or [],
            float(payload.get("roeToInr") or 1),
            flights,
            payload.get("revRows") or [],
            payload.get("revHeaders") or {},
        )
        if str(profile.get("layout") or "") != "maa":
            stamp_rev_summary_title(wb[rev], str(payload.get("from") or ""))
    if "Addl & XXLD Flt" in wb.sheetnames:
        populate_addl_xxld(wb["Addl & XXLD Flt"], flights, token, station_label)

    wb.save(output)


def main() -> int:
    if len(sys.argv) < 2:
        print("Usage: export-bhs-mis-openpyxl.py <payload.json>", file=sys.stderr)
        return 1
    payload = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8-sig"))
    export(payload)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
