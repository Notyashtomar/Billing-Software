#!/usr/bin/env python3
"""Populate ADS MIS template with ground equipment & billing data — preserves Excel formulas (openpyxl)."""
from __future__ import annotations

import json
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

AIRLINE_FONT = Font(bold=True, name="Calibri", size=11)
EQUIP_FONT = Font(bold=False, name="Calibri", size=10)
HEADER_FONT = Font(bold=True, name="Calibri", size=11)
HEADER_ALIGN = Alignment(wrap_text=False, vertical="center", horizontal="left")

ADS_BILLING_HEADER_ROW = 2
ADS_BILLING_FIRST_ROW = 3
ADS_BILLING_COL = {
    "label": 1,
    "used": 2,
    "chargeable": 3,
    "revenue": 4,
    "royalty": 5,
    "gst": 6,
    "total": 7,
    "remark": 8,
    "invoiceValue": 9,
    "invoiceNo": 10,
}
ADS_WORKING_FIRST_ROW = 2
ADS_WORKING_COL = {
    "serial": 1,
    "customerName": 2,
    "customerCode": 3,
    "flightNo": 4,
    "stdDate": 5,
    "atdDate": 6,
    "equipmentNo": 7,
    "unitType": 8,
    "quantity": 9,
}


ADS_MISSING = "-"


def cell_value(value: Any) -> Any:
    if value is None:
        return ADS_MISSING
    if isinstance(value, str) and not value.strip():
        return ADS_MISSING
    if isinstance(value, float) and (value != value):  # NaN
        return ADS_MISSING
    return value


def header_column_width(header: str) -> float:
    n = len((header or "").strip())
    return max(12, min(42, n + 4))


def fit_columns_to_headers(ws: Worksheet, headers: tuple[str, ...] | list[str], header_row: int) -> None:
    for i, title in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=i)
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        ws.column_dimensions[get_column_letter(i)].width = header_column_width(str(title))


def set_rc(ws: Worksheet, row: int, col: int, value: Any) -> None:
    cell = ws.cell(row=row, column=col)
    if type(cell).__name__ == "MergedCell":
        return
    cell.value = cell_value(value)


def month_token(from_ymd: str) -> str:
    m = re.match(r"^(\d{4})-(\d{2})", from_ymd or "")
    if not m:
        return "MON-YY"
    abbr = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]
    return f"{abbr[int(m.group(2)) - 1]}-{m.group(1)[2:]}"


def month_display(from_ymd: str) -> str:
    m = re.match(r"^(\d{4})-(\d{2})", from_ymd or "")
    if not m:
        return ""
    names = [
        "January", "February", "March", "April", "May", "June",
        "July", "August", "September", "October", "November", "December",
    ]
    return f"{names[int(m.group(2)) - 1]}-{m.group(1)[2:]}"


def replace_month_tokens(ws: Worksheet, token: str, station_label: str, display: str) -> None:
    max_row = min(ws.max_row or 0, 10)
    max_col = min(ws.max_column or 0, 25)
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(r, c)
            if cell.data_type == "f" or not isinstance(cell.value, str):
                continue
            text = cell.value
            text = re.sub(r"[A-Z]{3}-\d{2}", token, text)
            text = re.sub(
                r"for the month of [A-Za-z]+-\d{2}",
                f"for the month of {token}",
                text,
                flags=re.I,
            )
            text = re.sub(
                r"for the m/o [A-Za-z]+-\d{2}",
                f"for the m/o {token}",
                text,
                flags=re.I,
            )
            text = re.sub(r"Chennai Stat\w*", f"{station_label} Station", text, flags=re.I)
            if display:
                text = re.sub(
                    r"(ADS WORKING FOR THE MONTH OF )[A-Za-z]+-\d{2}( FOR )\w+",
                    rf"\1{display}\2{station_label}",
                    text,
                    flags=re.I,
                )
            if text != cell.value:
                cell.value = text


def populate_working_sheet(ws: Worksheet, eq_rows: list[dict[str, Any]]) -> None:
    start_row = ADS_WORKING_FIRST_ROW
    col = ADS_WORKING_COL

    for idx, row in enumerate(eq_rows):
        r = start_row + idx
        set_rc(ws, r, col["serial"], row.get("serialNumber") or (idx + 1))
        set_rc(ws, r, col["customerName"], row.get("customerName"))
        set_rc(ws, r, col["customerCode"], row.get("customerCode"))
        set_rc(ws, r, col["flightNo"], row.get("flightNo"))
        set_rc(ws, r, col["stdDate"], row.get("stdDate"))
        set_rc(ws, r, col["atdDate"], row.get("atdDate"))
        set_rc(ws, r, col["equipmentNo"], row.get("equipmentNo"))
        set_rc(ws, r, col["unitType"], row.get("unitType") or "Unit")
        set_rc(ws, r, col["quantity"], row.get("quantity"))


def populate_exported_eq_data(ws: Worksheet, eq_rows: list[dict[str, Any]]) -> None:
    populate_working_sheet(ws, eq_rows)


def populate_billing_sheet(ws: Worksheet, billing_rows: list[dict[str, Any]]) -> None:
    """Airline header (bold totals) + equipment lines. Cols match MAA July workbook."""
    headers = (
        "Row Labels",
        "Used units",
        "Chargeable units",
        "Revenue",
        "Royalty",
        "GST",
        "Total Invoice Amount",
        "Remark",
        "Invoice Value",
        "Invoice No",
    )
    for col, title in enumerate(headers, start=1):
        set_rc(ws, ADS_BILLING_HEADER_ROW, col, title)

    col = ADS_BILLING_COL
    for idx, row in enumerate(billing_rows):
        r = ADS_BILLING_FIRST_ROW + idx
        kind = str(row.get("kind") or "equipment")
        revenue_raw = row.get("revenue")
        invoices = str(row.get("invoiceNumbers") or "").strip()

        set_rc(ws, r, col["label"], row.get("label"))
        set_rc(ws, r, col["used"], row.get("used"))
        set_rc(ws, r, col["chargeable"], row.get("chargeable"))
        set_rc(ws, r, col["revenue"], row.get("revenue"))
        set_rc(ws, r, col["royalty"], row.get("royalty"))
        set_rc(ws, r, col["gst"], row.get("gst"))
        set_rc(ws, r, col["total"], row.get("total"))
        set_rc(ws, r, col["remark"], row.get("remark"))
        try:
            revenue = float(revenue_raw)
        except (TypeError, ValueError):
            revenue = 0.0
        set_rc(
            ws,
            r,
            col["invoiceValue"],
            row.get("total") if revenue > 0 or kind == "airline" else ADS_MISSING,
        )
        set_rc(ws, r, col["invoiceNo"], invoices if invoices else ADS_MISSING)

        font = AIRLINE_FONT if kind == "airline" else EQUIP_FONT
        for c in range(1, 11):
            cell = ws.cell(row=r, column=c)
            cell.font = font
            cell.alignment = Alignment(wrap_text=False, vertical="center")
        ws.row_dimensions[r].height = 15

    last = ADS_BILLING_FIRST_ROW + len(billing_rows)
    for r in range(last, last + 20):
        if ws.cell(row=r, column=1).value is not None:
            for c in range(1, 11):
                set_rc(ws, r, c, None)

    fit_columns_to_headers(ws, headers, ADS_BILLING_HEADER_ROW)


def process_ads_mis(payload: dict[str, Any]) -> None:
    template = Path(payload["templatePath"])
    out_path = Path(payload["outputPath"])

    token = month_token(payload.get("from") or "")
    display = month_display(payload.get("from") or "")
    station_label = payload.get("stationLabel") or "Chennai"

    wb = openpyxl.load_workbook(template)

    for ws in wb.worksheets:
        replace_month_tokens(ws, token, station_label, display)

    eq_rows = payload.get("equipmentRows") or []
    invoices = payload.get("invoices") or []
    billing_rows = payload.get("billingRows") or []

    if "WorkingSheet" in wb.sheetnames:
        populate_working_sheet(wb["WorkingSheet"], eq_rows)

    if "Exported EQ Data" in wb.sheetnames:
        populate_exported_eq_data(wb["Exported EQ Data"], eq_rows)

    if "Billing Sheet" in wb.sheetnames:
        populate_billing_sheet(
            wb["Billing Sheet"],
            billing_rows
            if billing_rows
            else [
                {
                    "kind": "airline",
                    "label": inv.get("customerName") or inv.get("customerCode"),
                    "used": 0,
                    "chargeable": 0,
                    "revenue": 0,
                    "royalty": 0,
                    "gst": 0,
                    "total": inv.get("totalInclTax") or 0,
                    "invoiceNumbers": inv.get("invoiceNumber") or "",
                }
                for inv in invoices
            ],
        )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    wb.close()


def main() -> None:
    if len(sys.argv) < 2:
        print("Usage: export-ads-mis-openpyxl.py <payload.json>", file=sys.stderr)
        sys.exit(1)

    payload_file = Path(sys.argv[1])
    if not payload_file.exists():
        print(f"Payload file not found: {payload_file}", file=sys.stderr)
        sys.exit(1)

    with open(payload_file, encoding="utf-8") as f:
        payload = json.load(f)

    process_ads_mis(payload)


if __name__ == "__main__":
    main()
